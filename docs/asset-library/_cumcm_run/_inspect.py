import openpyxl, glob, os

wd = r'D:\modex\_assets_extracted\_cumcm_run\workspace\user_data'
for f in sorted(glob.glob(os.path.join(wd, '*.xlsx'))):
    try:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        print('FILE:', os.path.basename(f))
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
            print('  SHEET:', ws.title, '| dims:', ws.max_row, 'x', ws.max_column)
            print('    hdr:', rows[0])
            if len(rows) > 1:
                print('    r2 :', rows[1])
        wb.close()
    except Exception as e:
        print('FILE:', os.path.basename(f), 'ERR:', e)

print('\n=== python 依赖检查 ===')
for m in ['numpy', 'scipy', 'pandas', 'openpyxl', 'mpmath', 'matplotlib']:
    try:
        __import__(m)
        print(m, 'OK')
    except Exception:
        print(m, 'MISSING')